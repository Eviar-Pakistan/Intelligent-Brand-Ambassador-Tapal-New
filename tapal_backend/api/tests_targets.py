from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from rest_framework.test import APIClient

from .models import Ambassador, AmbassadorMonthTarget, Store
from .shifts import build_monthly_shift, parse_hhmm
from .target_sheet import SKU_CATALOGUE

STORE = 'Al Latif Super Store (Multan)'


def sheet(rows):
    """Same layout as 'BA targets multan.xlsx'."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(['Region', 'City', 'Store name', 'Brand', 'SKU Name', 'KG Count', 'Count', 'Grammage'])
    for i, (store, brand, sku, kg, grams) in enumerate(rows, start=2):
        ws.append(['Multan', 'Multan', store, brand, sku, kg, f'=F{i}/H{i}', grams])
    out = BytesIO()
    wb.save(out)
    return SimpleUploadedFile('targets.xlsx', out.getvalue())


ROWS = [
    (STORE, 'Danedar', 'DD 100gm Tea Bag', 5, 0.1),
    (STORE, 'Danedar', 'DD 1750gm Pouch', 16.5, 1.65),
    (STORE, 'Tezdum', 'TD 290gm Pouch', None, 0),
    (STORE, 'Green Tea', 'Elaichi 45gm', 0.9, 0.045),
]


class StoreSkuTargetUploadTests(TestCase):
    def setUp(self):
        User = get_user_model()
        self.ho = APIClient()
        self.ho.force_authenticate(User.objects.create_user(username='ho', email='ho@x.com', password='pw', user_type=1))
        self.store = Store.objects.create(name=STORE, store_code='MUX-1', city='Multan', address='x')
        self.ba = Ambassador.objects.create(name='Zubair BA', city='Multan', status=Ambassador.Status.CERTIFIED)
        build_monthly_shift(
            store=self.store, ambassador=self.ba, month='2026-10', start=parse_hhmm('10:00'), end=parse_hhmm('18:00')
        ).save()

    def upload(self, client, rows=ROWS, month='2026-10'):
        return client.post('/api/ba-targets/upload/', {'file': sheet(rows), 'month': month}, format='multipart')

    def test_assigns_store_targets_to_its_ba(self):
        res = self.upload(self.ho)
        self.assertEqual(res.status_code, 201, res.data)
        self.assertEqual(res.data['saved'][0]['baName'], 'Zubair BA')
        target = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        self.assertEqual(float(target.target_total), 22.4)
        lines = {line['sku']: line for line in target.lines}
        self.assertEqual(lines['DD 100gm Tea Bag']['unit'], 50.0)       # 5 kg / 0.1 kg per pack
        self.assertEqual(lines['DD 1750gm Pouch']['unit'], 10.0)
        self.assertIsNone(lines['TD 290gm Pouch']['unit'])              # grammage 0 -> no #DIV/0!
        self.assertEqual(lines['Elaichi 45gm']['brand'], 'Green Tea')
        shown = self.ho.get('/api/ba-targets/?month=2026-10').data['results'][0]
        self.assertEqual((shown['storeName'], shown['targetKg'], len(shown['lines'])), (STORE, 22.4, 4))

    def test_reupload_replaces_targets_and_sales_come_from_reports(self):
        self.upload(self.ho)
        AmbassadorMonthTarget.objects.filter(ambassador=self.ba).update(sales_total=7)
        self.upload(self.ho, rows=[(STORE, 'Danedar', 'DD 100gm Tea Bag', 10, 0.1)])
        target = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        # no Daily Sales reports yet, so no sales
        self.assertEqual((float(target.target_total), float(target.sales_total), len(target.lines)), (10.0, 0.0, 1))

    def test_deployed_ba_also_gets_it(self):
        other = Ambassador.objects.create(name='Deployed BA', store=self.store, status=Ambassador.Status.DEPLOYED)
        res = self.upload(self.ho)
        self.assertEqual(sorted(s['baName'] for s in res.data['saved']), ['Deployed BA', 'Zubair BA'])
        self.assertTrue(AmbassadorMonthTarget.objects.filter(ambassador=other).exists())

    def test_reports_problems(self):
        res = self.upload(self.ho, rows=[('Unknown Store', 'Danedar', 'DD', 1, 0.1)])
        self.assertEqual(res.status_code, 400)
        self.assertEqual(res.data['unmatched_stores'], ['Unknown Store'])
        empty = self.upload(self.ho, rows=[(STORE, 'Danedar', 'DD', None, 0.1)])
        self.assertEqual(empty.data['stores_with_empty_kg'], [STORE])
        # a BA deployed at the store (store_id) gets its targets in any month; without that, no BA in November
        Ambassador.objects.filter(pk=self.ba.pk).update(store=None)
        no_ba = self.upload(self.ho, month='2026-11')
        self.assertEqual(no_ba.data['stores_without_ba'], [STORE])
        self.assertEqual(self.ho.post('/api/ba-targets/upload/', {'month': '2026-10'}).status_code, 400)

    def test_city_scope(self):
        User = get_user_model()
        lahore = APIClient()
        lahore.force_authenticate(
            User.objects.create_user(username='lhr', email='l@x.com', password='pw', user_type=1, city='Lahore')
        )
        res = self.upload(lahore)
        self.assertEqual(res.data['unmatched_stores'], [STORE])  # a Multan store is not theirs
        multan = APIClient()
        multan.force_authenticate(
            User.objects.create_user(username='mux', email='m@x.com', password='pw', user_type=1, city='Multan')
        )
        self.assertEqual(self.upload(multan).status_code, 201)


class StoreNameMatchingTests(TestCase):
    def test_hyphen_and_city_suffix(self):
        from .target_sheet import _find_store

        main = Store.objects.create(name='Al-Latif Super Store', store_code='A1', city='Multan', address='x')
        branch = Store.objects.create(name='Al-Latif Super Store -2', store_code='A2', city='Multan', address='x')
        stores = list(Store.objects.all())
        self.assertEqual(_find_store('Al Latif Super Store (Multan)', 'Multan', stores), main)
        self.assertEqual(_find_store('Al Latif Super Store -2 (Multan)', 'Multan', stores), branch)
        self.assertIsNone(_find_store('Al Latif Super Store (Lahore)', 'Lahore', stores[:0]))


class SkuTemplateTests(StoreSkuTargetUploadTests):
    def test_template_is_ba_code_by_sku(self):
        res = self.ho.get('/api/ba-targets/template/?month=2026-10')
        self.assertEqual(res.status_code, 200)
        book = openpyxl.load_workbook(BytesIO(res.content))
        rows = [r for r in book['Targets'].iter_rows(values_only=True) if any(r)]
        self.assertEqual(rows[0], ('BA Code', 'SKU', 'Month', 'Brand', 'Target Kg', 'Grammage', 'Target Units', 'Sales Kg', 'Sales Units'))
        self.assertTrue(str(rows[1][6]).startswith('=IF('))  # Units = Target Kg / Grammage formula
        self.assertEqual(len([r for r in book['SKU List'].iter_rows(values_only=True) if r[0]]) - 1, len(SKU_CATALOGUE))
        # saved targets come back filled in, one row per SKU
        self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'qty': 5}, {'sku': 'mint 45GM', 'qty': 1},
        ]}]}, format='json')
        res = self.ho.get('/api/ba-targets/template/?month=2026-10')
        rows = [r for r in openpyxl.load_workbook(BytesIO(res.content))['Targets'].iter_rows(values_only=True) if r[0]]
        self.assertEqual([r[:5] for r in rows[1:]], [
            (self.ba.ba_code, 'DD 100gm Tea Bag', '2026-10', 'Danedar', 5),
            (self.ba.ba_code, 'Mint 45gm', '2026-10', 'Green Tea', 1),
        ])

    def test_units_times_grammage(self):
        res = self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'units': 6}, {'sku': 'DD 900gm Pouch', 'units': 10},
        ]}]}, format='json')
        self.assertEqual(res.status_code, 201, res.data)
        t = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        lines = {l['sku']: (l['unit'], l['kg']) for l in t.lines}
        self.assertEqual(lines, {'DD 100gm Tea Bag': (6, 0.6), 'DD 900gm Pouch': (10, 9.0)})
        self.assertEqual(float(t.target_total), 9.6)

    def test_save_many_skus_per_ba(self):
        # a SKU that is not in the list: nothing saved for that BA
        bad = self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'qty': 5}, {'sku': 'Custom SKU', 'qty': 2},
        ]}]}, format='json')
        self.assertIn('Custom SKU — not in the SKU list', bad.data['errors'][0])
        self.assertFalse(AmbassadorMonthTarget.objects.filter(ambassador=self.ba, month='2026-10').exists())

        rows = [{'baCode': self.ba.ba_code.lower(), 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'qty': 5},
            {'sku': 'dd 900gm pouch', 'qty': 16.5, 'sales': 3},  # any case; sales in the upload are ignored
            {'sku': 'TD 80gm Hard Pack', 'brand': 'Wrong', 'qty': 2},  # brand comes from the SKU list
        ]}]
        res = self.ho.post('/api/ba-targets/', {'rows': rows}, format='json')
        self.assertEqual(res.status_code, 201, res.data)
        t = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        self.assertEqual((float(t.target_total), float(t.sales_total), t.store, len(t.lines)), (23.5, 0.0, self.store, 3))
        lines = {l['sku']: l for l in t.lines}
        self.assertEqual(lines['DD 100gm Tea Bag']['unit'], 50.0)
        self.assertEqual((lines['DD 900gm Pouch']['brand'], lines['TD 80gm Hard Pack']['brand']), ('Danedar', 'Tezdum'))
        self.assertEqual(lines['TD 80gm Hard Pack']['unit'], 25.0)
        shown = self.ho.get('/api/ba-targets/?month=2026-10').data['results'][0]
        self.assertEqual((shown['baCode'], len(shown['lines'])), (self.ba.ba_code, 3))
        # re-upload replaces the lines
        self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10',
                     'lines': [{'sku': 'DD 100gm Tea Bag', 'kg': 1}]}]}, format='json')
        t.refresh_from_db()
        self.assertEqual((float(t.target_total), float(t.sales_total), len(t.lines)), (1.0, 0.0, 1))

    def test_unit_sales_are_whole_numbers_and_count_as_kg(self):
        ba_app = APIClient()

        def report(report_id, sales):
            return ba_app.post('/api/daily-reports/', {
                'token': self.ba.invite_token, 'id': report_id, 'source': 'checkout',
                'submittedAt': '2026-10-03T18:00:00+05:00', 'stock': {}, 'sales': sales, 'otherBrands': [],
            }, format='json')

        self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'kg': 10},
        ]}]}, format='json')
        self.assertEqual(report('rep-dec', {'unit:DD 100gm Tea Bag': '2.5'}).status_code, 400)
        self.assertEqual(report('rep-ok', {'unit:DD 100gm Tea Bag': '12'}).status_code, 201)
        t = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        self.assertEqual(t.lines[0]['sales'], 1.2)  # 12 units x 0.1 kg
        self.assertEqual(float(t.sales_total), 1.2)

    def test_daily_sales_reports_drive_achievement(self):
        ba_app = APIClient()

        def report(report_id, when, sales):
            res = ba_app.post('/api/daily-reports/', {
                'token': self.ba.invite_token, 'id': report_id, 'source': 'checkout',
                'submittedAt': when, 'stock': {}, 'sales': sales, 'otherBrands': [],
            }, format='json')
            self.assertEqual(res.status_code, 201, res.data)

        # a report before any target exists still counts once the target is set
        report('rep-1', '2026-10-03T18:00:00+05:00', {'sku:DD 100gm Tea Bag': '2.5', 'totalInterceptions': '9'})
        self.ho.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [
            {'sku': 'DD 100gm Tea Bag', 'kg': 10}, {'sku': 'TD 80gm Hard Pack', 'qty': 4},
        ]}]}, format='json')
        # next day: two reports, the later one replaces the earlier one for that day
        report('rep-2', '2026-10-04T12:00:00+05:00', {'sku:DD 100gm Tea Bag': '1', 'sku:TD 80gm Hard Pack': '1'})
        report('rep-3', '2026-10-04T19:00:00+05:00', {'sku:DD 100gm Tea Bag': '1.5', 'sku:TD 80gm Hard Pack': '2'})
        # another month does not count
        report('rep-4', '2026-11-01T12:00:00+05:00', {'sku:DD 100gm Tea Bag': '50'})

        t = AmbassadorMonthTarget.objects.get(ambassador=self.ba, month='2026-10')
        lines = {l['sku']: l['sales'] for l in t.lines}
        self.assertEqual(lines, {'DD 100gm Tea Bag': 4.0, 'TD 80gm Hard Pack': 2.0})
        self.assertEqual((float(t.target_total), float(t.sales_total)), (14.0, 6.0))
        shown = self.ho.get('/api/ba-targets/?month=2026-10').data['results'][0]
        self.assertEqual(round(shown['salesKg'] / shown['targetKg'] * 100), 43)  # achievement %

    def test_ba_without_store_and_bad_code(self):
        loose = Ambassador.objects.create(name='No Store BA', city='Multan')
        res = self.ho.post('/api/ba-targets/', {'rows': [
            {'baCode': loose.ba_code, 'month': '2026-10', 'lines': [{'sku': 'TD 80gm Hard Pack', 'qty': 2}]},
            {'baCode': 'BA-NOPE', 'month': '2026-10', 'lines': [{'sku': 'TD 80gm Hard Pack', 'qty': 2}]},
        ]}, format='json')
        self.assertEqual(len(res.data['saved']), 1)
        self.assertIn('BA-NOPE', res.data['errors'][0])
        self.assertIsNone(AmbassadorMonthTarget.objects.get(ambassador=loose).store)
        self.assertEqual(self.ho.get('/api/ba-targets/?month=2026-10').data['results'][0]['storeName'], '')

    def test_city_user_only_their_bas(self):
        lahore = APIClient()
        lahore.force_authenticate(get_user_model().objects.create_user(username='l2', email='l2@x.com', password='pw', user_type=1, city='Lahore'))
        res = lahore.post('/api/ba-targets/', {'rows': [{'baCode': self.ba.ba_code, 'month': '2026-10', 'lines': [{'sku': 'TD 80gm Hard Pack', 'qty': 2}]}]}, format='json')
        self.assertEqual(res.status_code, 400)


class CityReportSkuTests(TestCase):
    def test_city_lists(self):
        from .target_sheet import canonical_sku, report_skus_for_city

        lahore, multan = report_skus_for_city('Lahore'), report_skus_for_city(' multan ')
        self.assertEqual((len(lahore), len(multan)), (41, 22))
        self.assertIsNone(report_skus_for_city('Faisalabad'))
        # rows in the City SKU table replace the built-in list for that city
        from .models import CitySku

        for name in ('dd 170gm hard pack', 'Danedar 900gm X 12 Pouch', 'New Pack 1kg'):
            CitySku.objects.create(sku=name, city='Faisalabad')
        self.assertEqual(report_skus_for_city('faisalabad'), [
            {'brand': 'Danedar', 'label': 'DD 170gm Hard Pack', 'sku': 'DD 170gm Hard Pack'},
            {'brand': 'Danedar', 'label': 'Danedar 900gm X 12 Pouch', 'sku': 'DD 900gm Pouch'},
            {'brand': 'Other', 'label': 'New Pack 1kg', 'sku': 'New Pack 1kg'},
        ])
        self.assertEqual(lahore[1], {'brand': 'Danedar', 'label': 'Danedar 170gm X 60 Hard Pack', 'sku': 'DD 170gm Hard Pack'})
        # every Multan SKU is on the Lahore list, and mapped SKUs are real target SKUs
        self.assertTrue({s['label'] for s in multan} <= {s['label'] for s in lahore})
        for item in lahore:
            if item['sku'] != item['label']:
                self.assertEqual(canonical_sku(item['sku']), item['sku'])
        self.assertEqual(len({s['sku'] for s in lahore}), 41)

    def test_ba_me_sends_the_city_list(self):
        ba = Ambassador.objects.create(name='Lahore BA', city='Lahore')
        other = Ambassador.objects.create(name='Fsd BA', city='Faisalabad')
        client = APIClient()
        self.assertEqual(len(client.get(f'/api/ba/me/?token={ba.invite_token}').data['reportSkus']), 41)
        self.assertIsNone(client.get(f'/api/ba/me/?token={other.invite_token}').data['reportSkus'])
